"""PDF (pypdf) and DOCX (python-docx) text extraction."""

import io
import zipfile
from typing import Final

from docx import Document
from pypdf import PdfReader

from autoapplier.domain.resume_files import DocumentKind
from autoapplier.ports.documents import DocumentUnreadableError

DOCX_MAX_UNCOMPRESSED_BYTES: Final = 50 * 1024 * 1024


class PyPdfDocxTextExtractor:
    """`DocumentTextExtractor` backed by pypdf and python-docx."""

    def __init__(self, *, max_docx_uncompressed_bytes: int = DOCX_MAX_UNCOMPRESSED_BYTES) -> None:
        self._max_docx_uncompressed_bytes = max_docx_uncompressed_bytes

    def extract_text(self, data: bytes, kind: DocumentKind) -> str:
        try:
            if kind == "pdf":
                return self._pdf(data)
            return self._docx(data)
        except DocumentUnreadableError:
            raise
        except Exception as exc:  # third-party parsers raise many types; all mean unreadable
            raise DocumentUnreadableError(f"cannot read {kind}: {type(exc).__name__}") from exc

    def _guard_docx_size(self, data: bytes) -> None:
        """Refuse a zip whose declared uncompressed size is a bomb, before anything inflates it."""
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
        if total > self._max_docx_uncompressed_bytes:
            raise DocumentUnreadableError("docx too large when uncompressed")

    @staticmethod
    def _pdf(data: bytes) -> str:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentUnreadableError("encrypted pdf")
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def _docx(self, data: bytes) -> str:
        self._guard_docx_size(data)
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
